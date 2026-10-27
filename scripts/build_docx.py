#!/usr/bin/env python3
"""Editable Word version of the guide, for clients who want to change the wording.

  build_docx.py --ws W        -> W/final/user-guide.docx

Same content as the PDF: cover, contents, modules, numbered steps with annotated
screenshots, tips, warnings, field tables, cheat sheets, common problems.
Needs python-docx. The contents list is a Word field; Word offers to update it when
the file opens (click Yes), or right-click it and choose Update field.
"""
import argparse, json, os, re, sys, time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: pip install --break-system-packages python-docx")
from PIL import Image
from build_pdf import before_lines, load_common, load_state, resolve_step, step_image
from diagram import flow_png

BOLD = re.compile(r"\*\*(.+?)\*\*")


def hex_to_rgb(h, default=(0xE8, 0x31, 0x26)):
    try:
        h = h.lstrip("#")
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return default


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def left_bar(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    b = OxmlElement("w:left")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "24")
    b.set(qn("w:color"), color)
    borders.append(b)
    tcPr.append(borders)


def add_runs(par, text, size=None, color=None):
    pos = 0
    for m in BOLD.finditer(text or ""):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            style_run(r, size, color)
        r = par.add_run(m.group(1))
        r.bold = True
        style_run(r, size, color)
        pos = m.end()
    if pos < len(text or ""):
        r = par.add_run(text[pos:])
        style_run(r, size, color)


def style_run(r, size, color):
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)


def add_field(par, instr, placeholder=""):
    for kind, extra in (("begin", None), (None, instr), ("separate", None), (None, placeholder), ("end", None)):
        run = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif extra is instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            run._r.append(it)
        else:
            run.text = placeholder


def callout(doc, label, text, fill, bar):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, fill)
    left_bar(c, bar)
    p = c.paragraphs[0]
    r = p.add_run(f"{label} ")
    r.bold = True
    r.font.size = Pt(10)
    add_runs(p, text, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def grid_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, "D6E7F1")
        c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].text = ""
            add_runs(cells[i].paragraphs[0], v, 10)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def add_picture(doc, path, max_w=6.2, max_h=5.0):
    with Image.open(path) as im:
        w, h = im.size
    wi = min(max_w, max_h * w / h)
    p = doc.add_paragraph()
    p.add_run().add_picture(path, width=Inches(wi))
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = False


def build(ws):
    state = load_state(ws)
    common = load_common(ws, state)
    ready = [s for s in state["order"] if state["modules"][s]["status"] == "built"]
    ready.sort(key=lambda s: {"quick-start": 0, "workflows": 1}.get(s, 2))
    if not ready:
        sys.exit("no built modules yet")
    br = state.get("brand") or {}
    accent = hex_to_rgb(br.get("accent") or "#1C75BC")
    accent_hex = "%02X%02X%02X" % accent

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(20))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for name, size in (("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    # cover
    logo = br.get("logo")
    if logo and os.path.exists(logo):
        with Image.open(logo) as li:
            lw, lh = li.size
        doc.add_paragraph().add_run().add_picture(logo, width=Inches(min(2.4, 0.9 * lw / lh)))
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph()
    r = t.add_run(state["app"])
    r.bold = True
    r.font.size = Pt(34)
    s = doc.add_paragraph()
    r = s.add_run("User guide")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(*accent)
    m = doc.add_paragraph(time.strftime("%d %B %Y") + (f"\n{state['url']}" if state.get("url") else ""))
    m.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    note = doc.add_paragraph("Numbered marks on the screenshots show where to click.")
    note.runs[0].font.size = Pt(10)
    doc.add_page_break()

    # contents (Word field)
    doc.add_heading("Contents", level=1)
    cp = doc.add_paragraph()
    add_field(cp, r'TOC \o "1-2" \h \z \u', "Right-click here and choose Update field to show the contents.")
    doc.add_page_break()

    for slug in ready:
        with open(os.path.join(ws, "modules", f"{slug}.json")) as f:
            mod = json.load(f)
        doc.add_heading(mod.get("title", slug), level=1)
        if mod.get("intro"):
            p = doc.add_paragraph()
            add_runs(p, mod["intro"], 11, (0x6B, 0x72, 0x80))
        if mod.get("cheatsheet"):
            grid_table(doc, ["I want to...", "Go to"],
                       [[r_.get("want", ""), r_.get("go", "")] for r_ in mod["cheatsheet"]], [3.1, 3.1])
        for feat in mod.get("features", []):
            doc.add_heading(feat["title"], level=2)
            if feat.get("goal"):
                p = doc.add_paragraph()
                add_runs(p, feat["goal"], 10.5, (0x6B, 0x72, 0x80))
            if feat.get("gif"):
                p = doc.add_paragraph()
                r = p.add_run(f"Animated version: {feat['gif']} (see final/guide.html)")
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            bl = before_lines(feat)
            if bl:
                callout(doc, "Before you start.", "  ".join(f"{a}: {b}." if not b.endswith(".") else f"{a}: {b}" for a, b in bl), "D6E7F1", accent_hex)
            for i, step in enumerate((resolve_step(x, common, ws, state) for x in feat.get("steps", [])), 1):
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(6)
                num = p.add_run(f"{i}  ")
                num.bold = True
                num.font.color.rgb = RGBColor(*accent)
                num.font.size = Pt(13)
                add_runs(p, step["text"], 11)
                img = step_image(step, ws)
                if img:
                    add_picture(doc, img)
                if step.get("warning"):
                    callout(doc, "Careful.", step["warning"], "FDECEA", "C0392B")
                if step.get("tip"):
                    callout(doc, "Tip.", step["tip"], "D6E7F1", accent_hex)
            if feat.get("fields"):
                doc.add_heading("Fields on this screen", level=3)
                grid_table(doc, ["Field", "Required?", "What to enter", "Example"],
                           [[f_.get("name", ""), "Required" if f_.get("required") else "Optional",
                             f_.get("help", ""), f_.get("example", "")] for f_ in feat["fields"]],
                           [1.5, 0.9, 2.5, 1.3])
        for wf in mod.get("workflows") or []:
            doc.add_heading(wf["title"], level=2)
            if wf.get("goal"):
                add_runs(doc.add_paragraph(), wf["goal"], 10.5, (0x6B, 0x72, 0x80))
            add_picture(doc, flow_png(ws, wf, "#%02X%02X%02X" % accent), max_w=6.3, max_h=4.2)
            grid_table(doc, ["#", "What to do", "Where", "Note"],
                       [[str(i), f_.get("label", ""), f_.get("where", ""), f_.get("note", "")]
                        for i, f_ in enumerate(wf.get("flow", []), 1)], [0.4, 2.0, 2.0, 1.8])
        if mod.get("problems"):
            doc.add_heading("Common problems", level=2)
            grid_table(doc, ["You see", "Why it happens", "What to do"],
                       [[p_.get("message", ""), p_.get("cause", ""), p_.get("fix", "")] for p_ in mod["problems"]],
                       [2.1, 1.9, 2.2])
        doc.add_page_break()

    # footer page numbers
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", "1")

    # ask Word to refresh fields (contents) on open
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)

    os.makedirs(os.path.join(ws, "final"), exist_ok=True)
    out = os.path.join(ws, "final", "user-guide.docx")
    doc.save(out)
    print(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--ws", required=True)
    build(ap.parse_args().ws)
